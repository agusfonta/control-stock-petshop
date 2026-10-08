"""Lectura de planillas xlsx/csv a filas crudas (C-08, D10).

Pura: recibe bytes y devuelve una `Planilla` (nunca toca la base). Los
problemas del contenido (archivo roto, hoja inexistente, demasiadas filas) son
errores globales en `Planilla.errores`; solo un formato desconocido levanta.
"""

import csv
import io
from dataclasses import dataclass, field
from pathlib import PurePath

from openpyxl import load_workbook

from app.services.migracion.normalizar import clave

MAX_FILAS = 5000
EXTENSIONES = (".xlsx", ".csv")


class ArchivoNoAdmitido(Exception):
    """Extension distinta de .xlsx/.csv (se traduce a 415)."""


@dataclass(frozen=True)
class Celda:
    valor: object
    es_porcentaje: bool = False
    formula_sin_valor: bool = False


@dataclass(frozen=True)
class FilaCruda:
    numero: int  # fila real del Excel (1-based, el encabezado cuenta)
    celdas: tuple[Celda, ...]


@dataclass(frozen=True)
class Planilla:
    encabezados: tuple[object, ...] = ()
    filas: tuple[FilaCruda, ...] = ()
    errores: tuple[str, ...] = field(default_factory=tuple)


def _con_error(mensaje: str) -> Planilla:
    return Planilla(errores=(mensaje,))


def _armar(filas_crudas: list[tuple[int, list[Celda]]]) -> Planilla:
    """Primera fila no vacia = encabezado; el resto, datos sin filas vacias."""
    no_vacias = [
        (n, celdas)
        for n, celdas in filas_crudas
        if any(c.valor not in (None, "") for c in celdas)
    ]
    if not no_vacias:
        return _con_error("la planilla esta vacia: no tiene encabezados")
    _, enc = no_vacias[0]
    encabezados = tuple(c.valor for c in enc)
    datos = no_vacias[1:]
    if len(datos) > MAX_FILAS:
        return _con_error(
            f"la planilla supera el maximo de {MAX_FILAS} filas de datos "
            f"({len(datos)}); dividila en archivos mas chicos"
        )
    ancho = len(encabezados)
    vacia = Celda(None)
    filas = tuple(
        FilaCruda(n, tuple((celdas + [vacia] * ancho)[:ancho])) for n, celdas in datos
    )
    return Planilla(encabezados, filas)


def _leer_csv(contenido: bytes) -> Planilla:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("cp1252", errors="replace")
    primera = next((ln for ln in texto.splitlines() if ln.strip()), "")
    delimitador = ";" if primera.count(";") > primera.count(",") else ","
    filas = [
        (n, [Celda(v.strip() if v.strip() else None) for v in registro])
        for n, registro in enumerate(csv.reader(io.StringIO(texto), delimiter=delimitador), 1)
    ]
    return _armar(filas)


def _formulas(contenido: bytes, hoja: str) -> set[tuple[int, int]]:
    """Posiciones (fila, columna) 1-based de celdas con formula."""
    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
    try:
        ws = wb[hoja]
        ws.reset_dimensions()  # algunos generadores declaran mal el rango
        return {
            (c.row, c.column)
            for fila in ws.iter_rows()
            for c in fila
            if getattr(c, "data_type", None) == "f"
        }
    finally:
        wb.close()


def _leer_xlsx(contenido: bytes, hoja: str | None) -> Planilla:
    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        if hoja is None:
            nombre = wb.sheetnames[0]
        else:
            nombre = next((s for s in wb.sheetnames if clave(s) == clave(hoja)), None)
            if nombre is None:
                return _con_error(
                    f"no existe la hoja '{hoja}'; hojas disponibles: "
                    + ", ".join(wb.sheetnames)
                )
        ws = wb[nombre]
        ws.reset_dimensions()  # algunos generadores declaran mal el rango
        crudas = [
            [(c.value, "%" in (getattr(c, "number_format", "") or "")) for c in fila]
            for fila in ws.iter_rows()
        ]
    finally:
        wb.close()
    hay_huecos = any(v is None for fila in crudas for v, _ in fila)
    formulas = _formulas(contenido, nombre) if hay_huecos else set()
    filas = [
        (
            n,
            [
                Celda(v, es_porcentaje=pct and isinstance(v, (int, float)),
                      formula_sin_valor=v is None and (n, col) in formulas)
                for col, (v, pct) in enumerate(fila, start=1)
            ],
        )
        for n, fila in enumerate(crudas, start=1)
    ]
    return _armar(filas)


def leer_planilla(contenido: bytes, nombre_archivo: str, hoja: str | None = None) -> Planilla:
    """Lee un .xlsx o .csv. Levanta ArchivoNoAdmitido por otra extension."""
    extension = PurePath(nombre_archivo).suffix.lower()
    if extension not in EXTENSIONES:
        raise ArchivoNoAdmitido(
            f"formato no admitido ({extension or 'sin extension'}): use .xlsx o .csv"
        )
    if extension == ".csv":
        return _leer_csv(contenido)
    try:
        return _leer_xlsx(contenido, hoja)
    except Exception:  # noqa: BLE001 - openpyxl levanta tipos muy variados
        return _con_error(
            "no se pudo leer el archivo xlsx: puede estar danado; "
            "abrilo y guardalo de nuevo en Excel"
        )
