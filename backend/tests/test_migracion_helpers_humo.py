"""Humo de los helpers de planillas de C-08 (grupo 1.3)."""

from io import BytesIO

from openpyxl import load_workbook

from tests.migracion_helpers import csv_bytes, xlsx_bytes


def test_xlsx_bytes_roundtrip_conserva_celdas():
    contenido = xlsx_bytes(
        ["sku", "costo"], [["A-1", 100], ["A-2", "$ 5,50"]], hoja="Productos"
    )
    ws = load_workbook(BytesIO(contenido)).active
    assert ws.title == "Productos"
    assert [[c.value for c in fila] for fila in ws.iter_rows()] == [
        ["sku", "costo"],
        ["A-1", 100],
        ["A-2", "$ 5,50"],
    ]


def test_xlsx_bytes_aplica_formatos_por_celda_y_formulas():
    contenido = xlsx_bytes(
        ["sku", "margen", "costo"],
        [["A-1", 0.35, "=10*2"]],
        formatos={(2, 2): "0%"},
    )
    ws = load_workbook(BytesIO(contenido)).active
    assert ws.cell(2, 2).number_format == "0%"
    assert ws.cell(2, 3).value == "=10*2"


def test_csv_bytes_codifica_con_el_encoding_pedido():
    assert csv_bytes("a;b\nÑandú;1", "cp1252") == "a;b\nÑandú;1".encode("cp1252")
    assert csv_bytes("x", "utf-8-sig").startswith(b"\xef\xbb\xbf")
