"""Generador determinista de planillas de prueba (C-08, tarea 9.1)."""

import pytest
from openpyxl import load_workbook

from app.models import Usuario
from app.services.migracion import importar
from scripts.generar_excel_prueba import generar

LIMPIA = "productos_prueba.xlsx"
SUCIA = "productos_prueba_sucio.xlsx"


@pytest.fixture
def db(db_session_factory):
    with db_session_factory() as session:
        yield session


@pytest.fixture
def duena(db):
    u = Usuario(email="d@test.only", password_hash="x", rol="duena")
    db.add(u)
    db.commit()
    return u


def celdas(ruta):
    wb = load_workbook(ruta)
    try:
        ws = wb.active
        return [
            [(c.value, c.number_format) for c in fila] for fila in ws.iter_rows()
        ]
    finally:
        wb.close()


def test_genera_las_dos_planillas_en_la_carpeta_pedida(tmp_path):
    destino = tmp_path / "ejemplos"
    generar(destino)
    assert sorted(p.name for p in destino.iterdir()) == [LIMPIA, SUCIA]


def test_la_planilla_limpia_tiene_volumen_y_encabezados_canonicos(tmp_path):
    generar(tmp_path)
    filas = celdas(tmp_path / LIMPIA)
    encabezados = [v for v, _ in filas[0]]
    assert encabezados[:3] == ["sku", "nombre", "categoria"]
    assert {"distribuidora", "costo", "margen_pct", "stock_inicial"} <= set(encabezados)
    datos = [{h: v for h, (v, _) in zip(encabezados, f)} for f in filas[1:]]
    assert len(datos) >= 15
    assert len({d["distribuidora"] for d in datos}) >= 2
    assert len({d["categoria"] for d in datos}) >= 2


def test_la_planilla_sucia_usa_alias_y_formatos_desordenados(tmp_path):
    generar(tmp_path)
    filas = celdas(tmp_path / SUCIA)
    encabezados = [v for v, _ in filas[0]]
    assert {"Código", "Proveedor", "Precio Costo", "Ganancia %"} <= set(encabezados)
    todo = [v for fila in filas[1:] for v, _ in fila]
    assert any(isinstance(v, str) and v.strip().startswith("$") for v in todo)
    assert any(isinstance(v, str) and v.endswith("%") for v in todo)
    assert any(isinstance(v, str) and v != v.strip() for v in todo)  # espacios sobrantes
    assert any(all(v is None for v, _ in fila) for fila in filas[1:])  # fila vacia
    assert any("%" in fmt for fila in filas[1:] for _, fmt in fila)  # margen con formato %


def test_dos_ejecuciones_producen_exactamente_las_mismas_celdas(tmp_path):
    generar(tmp_path / "a")
    generar(tmp_path / "b")
    for nombre in (LIMPIA, SUCIA):
        assert celdas(tmp_path / "a" / nombre) == celdas(tmp_path / "b" / nombre)


def test_analizar_la_limpia_da_todas_ok_y_crear(tmp_path, db, duena):
    generar(tmp_path)
    contenido = (tmp_path / LIMPIA).read_bytes()
    reporte = importar(db, contenido, LIMPIA, duena, confirmar=False)
    assert reporte.errores_globales == []
    assert reporte.totales.filas >= 15
    assert {(f.estado, f.accion) for f in reporte.filas} == {("ok", "crear")}
    assert reporte.totales.distribuidoras_a_crear >= 2


def test_analizar_la_sucia_muestra_ok_advertencias_y_los_tres_errores(tmp_path, db, duena):
    generar(tmp_path)
    contenido = (tmp_path / SUCIA).read_bytes()
    reporte = importar(db, contenido, SUCIA, duena, confirmar=False)
    assert reporte.errores_globales == []
    assert reporte.totales.ok >= 1 and reporte.totales.advertencias >= 1
    errores = {
        (m.campo, f.fila)
        for f in reporte.filas
        for m in f.motivos
        if m.gravedad == "error"
    }
    campos = {c for c, _ in errores}
    assert {"sku", "costo", "stock_inicial"} <= campos
    assert reporte.columnas.ignoradas  # hay una columna desconocida (Observaciones)
