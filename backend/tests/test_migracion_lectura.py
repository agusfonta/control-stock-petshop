"""Lectura de planillas xlsx/csv a filas crudas (C-08, D10)."""

from io import BytesIO

import pytest
from openpyxl import Workbook

from app.services.migracion.lectura import MAX_FILAS, ArchivoNoAdmitido, leer_planilla
from tests.migracion_helpers import csv_bytes, xlsx_bytes

ENC = ["sku", "nombre", "costo"]


def _valores(planilla):
    return [(f.numero, [c.valor for c in f.celdas]) for f in planilla.filas]


def test_xlsx_limpio_trae_encabezados_y_filas_con_numero_real():
    contenido = xlsx_bytes(ENC, [["A-1", "Uno", 10], ["A-2", "Dos", 20.5]])
    planilla = leer_planilla(contenido, "productos.xlsx")
    assert planilla.errores == ()
    assert list(planilla.encabezados) == ENC
    assert _valores(planilla) == [(2, ["A-1", "Uno", 10]), (3, ["A-2", "Dos", 20.5])]


def test_filas_vacias_intercaladas_se_saltean_conservando_la_numeracion():
    contenido = xlsx_bytes(
        ENC, [["A-1", "Uno", 1], [None, None, None], [None, None, None], ["A-4", "Cuatro", 4]]
    )
    planilla = leer_planilla(contenido, "p.xlsx")
    assert [f.numero for f in planilla.filas] == [2, 5]


def test_encabezado_precedido_de_filas_vacias_es_la_primera_fila_no_vacia():
    wb = Workbook()
    ws = wb.active
    for col, valor in enumerate(ENC, start=1):
        ws.cell(3, col, valor)
    for col, valor in enumerate(["A-1", "Uno", 5], start=1):
        ws.cell(4, col, valor)
    salida = BytesIO()
    wb.save(salida)
    planilla = leer_planilla(salida.getvalue(), "p.xlsx")
    assert list(planilla.encabezados) == ENC
    assert _valores(planilla) == [(4, ["A-1", "Uno", 5])]


def test_celda_con_formato_porcentaje_se_marca():
    contenido = xlsx_bytes([*ENC, "margen"], [["A-1", "Uno", 10, 0.35]], formatos={(2, 4): "0%"})
    celdas = leer_planilla(contenido, "p.xlsx").filas[0].celdas
    assert [c.es_porcentaje for c in celdas] == [False, False, False, True]


def test_formula_sin_valor_guardado_se_marca():
    contenido = xlsx_bytes(ENC, [["A-1", "Uno", "=10*2"]])
    celdas = leer_planilla(contenido, "p.xlsx").filas[0].celdas
    assert celdas[2].formula_sin_valor is True
    assert celdas[2].valor is None
    assert celdas[0].formula_sin_valor is False


def test_hoja_explicita_lee_esa_hoja():
    wb = Workbook()
    wb.active.title = "Otra"
    wb.active.append(["x"])
    stock = wb.create_sheet("Stock")
    stock.append(ENC)
    stock.append(["B-1", "Bolsa", 3])
    salida = BytesIO()
    wb.save(salida)
    planilla = leer_planilla(salida.getvalue(), "p.xlsx", hoja="Stock")
    assert _valores(planilla) == [(2, ["B-1", "Bolsa", 3])]


def test_hoja_inexistente_es_error_global():
    planilla = leer_planilla(xlsx_bytes(ENC, []), "p.xlsx", hoja="Nada")
    assert planilla.filas == ()
    assert any("Nada" in e for e in planilla.errores)


def test_csv_coma_utf8_con_bom():
    texto = "sku,nombre,costo\nA-1,Uno,10\n"
    planilla = leer_planilla(csv_bytes(texto, "utf-8-sig"), "p.csv")
    assert list(planilla.encabezados) == ENC
    assert _valores(planilla) == [(2, ["A-1", "Uno", "10"])]


def test_csv_punto_y_coma_cp1252_conserva_acentos():
    texto = "sku;nombre;costo\nA-1;Ñandú;10,5\n\nA-3;Perro;3\n"
    planilla = leer_planilla(csv_bytes(texto, "cp1252"), "p.CSV")
    assert _valores(planilla) == [(2, ["A-1", "Ñandú", "10,5"]), (4, ["A-3", "Perro", "3"])]


def test_csv_celdas_vacias_son_none():
    planilla = leer_planilla(csv_bytes("sku,nombre,costo\nA-1,,10\n"), "p.csv")
    assert planilla.filas[0].celdas[1].valor is None


@pytest.mark.parametrize("nombre", ["viejo.xls", "doc.pdf", "macro.xlsm", "sin_extension"])
def test_formatos_no_admitidos(nombre):
    with pytest.raises(ArchivoNoAdmitido):
        leer_planilla(b"lo que sea", nombre)


def test_xlsx_corrupto_es_error_global_legible():
    planilla = leer_planilla(b"esto no es un zip", "roto.xlsx")
    assert planilla.filas == ()
    assert len(planilla.errores) == 1
    assert "xlsx" in planilla.errores[0]


def test_planilla_sin_filas_es_error_global():
    planilla = leer_planilla(csv_bytes("\n\n"), "vacia.csv")
    assert planilla.errores and planilla.filas == ()


def test_limite_de_filas_de_datos_es_error_global():
    filas = [[f"S-{i}", "N", 1] for i in range(MAX_FILAS + 1)]
    planilla = leer_planilla(xlsx_bytes(ENC, filas), "grande.xlsx")
    assert planilla.filas == ()
    assert any(str(MAX_FILAS) in e for e in planilla.errores)


def test_justo_en_el_limite_se_acepta():
    filas = [[f"S-{i}", "N", 1] for i in range(MAX_FILAS)]
    planilla = leer_planilla(xlsx_bytes(ENC, filas), "limite.xlsx")
    assert planilla.errores == () and len(planilla.filas) == MAX_FILAS
